"""Build the Gridlock pitch deck (.pptx) — DARK THEME, full-height layouts.

Run: /opt/anaconda3/bin/python src/make_ppt.py  ->  Gridlock_Pitch_Deck.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import config as C

# ---- dark palette
BG   = RGBColor(0x0E, 0x13, 0x20)   # deep navy-black slide background
CARD = RGBColor(0x1B, 0x24, 0x36)   # panel
TXT  = RGBColor(0xF2, 0xF5, 0xFA)   # near-white
MUT  = RGBColor(0x9A, 0xA7, 0xB8)   # muted grey
LINE = RGBColor(0x2A, 0x35, 0x50)   # dividers / footer
BLUE = RGBColor(0x28, 0x74, 0xF0)   # fill (white text on it)
ABLUE= RGBColor(0x5A, 0xA0, 0xFF)   # accent text on dark
YEL  = RGBColor(0xFF, 0xE1, 0x1B)
GRN  = RGBColor(0x2E, 0xCC, 0x71)
AMB  = RGBColor(0xF4, 0xB7, 0x40)
ORG  = RGBColor(0xE6, 0x7E, 0x22)
RED  = RGBColor(0xFF, 0x5A, 0x5F)
FIG = C.FIG_DIR / "ppt"
FONT = "Calibri"

TEAM_NAME = "Team TrafficSolvers"
TEAM_MEMBERS = "Ryan Bhan (Lead, IIIT-D) · Kanan Mittal (IGDTUW) · Aditya Rawat (IIIT-D) · Hrithik Raj (IIIT-D)"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
_n = 0


def slide(dark_bg=True):
    s = prs.slides.add_slide(BLANK)
    if dark_bg:
        rect(s, 0, 0, W, H, BG)
    return s


def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=5):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space); p.space_before = Pt(0)
        for (txt, size, color, bold) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = FONT
    return tb


def header(s, kicker, title):
    rect(s, 0, Inches(0.32), Inches(0.16), Inches(0.95), YEL)
    text(s, Inches(0.5), Inches(0.30), Inches(12.3), Inches(0.35), [[(kicker, 12.5, ABLUE, True)]])
    text(s, Inches(0.5), Inches(0.60), Inches(12.4), Inches(0.7), [[(title, 27, TXT, True)]])
    rect(s, Inches(0.5), Inches(1.34), Inches(12.33), Pt(1.5), LINE)


def footer(s):
    global _n; _n += 1
    rect(s, Inches(0.5), Inches(7.0), Inches(12.33), Pt(1), LINE)
    text(s, Inches(0.5), Inches(7.06), Inches(8), Inches(0.35),
         [[("GRIDLOCK · Event-Driven Congestion Intelligence", 9.5, MUT, False)]])
    text(s, Inches(9.0), Inches(7.06), Inches(3.83), Inches(0.35),
         [[(f"Flipkart Gridlock 2.0 · {_n:02d}", 9.5, MUT, False)]], align=PP_ALIGN.RIGHT)


def pic_fit(s, name, bx, by, bw, bh):
    """Scale image to fit box (keep aspect) and center it within the box."""
    iw, ih = Image.open(FIG / name).size
    scale = min(bw / iw, bh / ih)
    w, h = int(iw * scale), int(ih * scale)
    x = bx + (bw - w) // 2; y = by + (bh - h) // 2
    s.shapes.add_picture(str(FIG / name), Emu(int(x)), Emu(int(y)), width=Emu(w), height=Emu(h))


def card(s, x, y, w, h, accent, title, body):
    rect(s, x, y, w, h, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, y, w, Inches(0.11), accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + Inches(0.2), y + Inches(0.28), w - Inches(0.4), h - Inches(0.5),
         [[(title, 15.5, TXT, True)], [(body, 12.5, MUT, False)]], space=6)


CB = Inches(6.78)   # content bottom guide

# =========================================================== 1 TITLE
s = slide()
rect(s, 0, 0, W, Inches(0.22), YEL)
rect(s, 0, Inches(7.28), W, Inches(0.22), BLUE)
text(s, Inches(0.9), Inches(2.05), Inches(11.5), Inches(1.2), [[("🚦 GRIDLOCK", 56, TXT, True)]])
text(s, Inches(0.95), Inches(3.35), Inches(11.6), Inches(0.7),
     [[("Event-Driven Congestion Intelligence for Bengaluru Traffic", 24, YEL, True)]])
text(s, Inches(0.95), Inches(4.2), Inches(11.6), Inches(0.6),
     [[("Forecast  →  Quantify Impact  →  Prescribe Deployment  →  Learn", 18, TXT, False)]])
rect(s, Inches(0.95), Inches(5.2), Inches(4.2), Pt(2), LINE)
text(s, Inches(0.95), Inches(5.4), Inches(11.8), Inches(1.7),
     [[(TEAM_NAME, 16, YEL, True)],
      [(TEAM_MEMBERS, 13, TXT, False)],
      [("Flipkart Gridlock Hackathon 2.0 · Round 2", 12.5, MUT, False)],
      [("Built on the anonymized ASTraM incident dataset — 8,173 events × 46 fields", 11.5, MUT, False)]])

# =========================================================== 2 PROBLEM
s = slide(); header(s, "THE OPERATIONAL CHALLENGE", "Events break the city — and we react blind")
text(s, Inches(0.5), Inches(1.5), Inches(12.3), Inches(0.9),
     [[("Rallies, festivals, sports, construction & sudden gatherings cause localized traffic "
        "breakdowns. Today the response is reactive and experience-driven.", 16, TXT, False)]])
cy, cw, ch = Inches(2.55), Inches(3.95), Inches(4.0)
card(s, Inches(0.5), cy, cw, ch, RED, "①  Impact is not quantified",
     "No way to know in advance how bad an event will be, how long it will tie up the road, "
     "or which junctions it threatens.")
card(s, Inches(4.69), cy, cw, ch, AMB, "②  Deployment is guesswork",
     "Officers, barricades and diversions are assigned from experience — not from data, and "
     "not optimized against a finite resource pool.")
card(s, Inches(8.88), cy, cw, ch, BLUE, "③  No post-event learning",
     "Every event is handled fresh. Mistakes are never fed back; the system never gets "
     "smarter after the fact.")
footer(s)

# =========================================================== 3 HOOK
s = slide(); header(s, "DATA INSIGHT MOST TEAMS WILL MISS", "This is an incident log — not congestion data")
pic_fit(s, "events_by_hour.png", Inches(0.4), Inches(1.55), Inches(7.8), Inches(5.2))
text(s, Inches(8.45), Inches(1.55), Inches(4.4), Inches(5.2),
     [[("There is NO speed / volume / delay field.", 16, RED, True)],
      [("So the target — “traffic impact” — does not exist in the data. It must be engineered.", 14, TXT, False)],
      [("", 8, TXT, False)],
      [("Events peak at 2 AM, not rush hour:", 16, YEL, True)],
      [("60% are vehicle breakdowns, driven by night-time freight movement. Teams that assume "
        "normal rush-hour patterns build the wrong model.", 14, TXT, False)],
      [("", 8, TXT, False)],
      [("We turn this gap into our centrepiece.", 15, ABLUE, True)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s)

# =========================================================== 4 REFRAME
s = slide(); header(s, "OUR REFRAME", "Not prediction — a decision-support system")
text(s, Inches(0.5), Inches(1.55), Inches(12.3), Inches(0.6),
     [[("Most teams stop at “predict congestion.” We deliver the full operational loop:", 16, TXT, False)]])
labels = [("FORECAST", "where / when / how many", BLUE),
          ("QUANTIFY", "Event Impact Score", GRN),
          ("PRESCRIBE", "manpower · barricades · diversions", AMB),
          ("LEARN", "post-event playbook", RED)]
x, bw, gap = Inches(0.6), Inches(2.85), Inches(0.25)
for i, (t1, t2, c) in enumerate(labels):
    rect(s, x, Inches(2.7), bw, Inches(2.5), c, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, Inches(3.25), bw, Inches(0.8), [[(t1, 21, TXT, True)]], align=PP_ALIGN.CENTER)
    text(s, x + Inches(0.12), Inches(4.2), bw - Inches(0.24), Inches(0.9),
         [[(t2, 13, TXT, False)]], align=PP_ALIGN.CENTER)
    if i < 3:
        text(s, x + bw, Inches(3.5), gap, Inches(1), [[("→", 26, MUT, True)]], align=PP_ALIGN.CENTER)
    x = x + bw + gap
text(s, Inches(0.6), Inches(5.75), Inches(12), Inches(0.9),
     [[("Two pipelines — Planned (proactive scheduling) and Unplanned (real-time triage) — "
        "exactly matching the problem statement.", 15, MUT, False)]])
footer(s)

# =========================================================== 5 ARCHITECTURE
s = slide(); header(s, "SYSTEM ARCHITECTURE", "Four layers on one data foundation")
rows = [("L0 · DATA FOUNDATION", "clean · leakage-safe features · H3 hex bins · EN+Kannada text", LINE),
        ("L1 · FORECAST", "LightGBM-Poisson corridor counts · empirical-Bayes hotspots · Hawkes self-excitation", BLUE),
        ("L2 · IMPACT (EIS)", "clearance quantile regression + road-closure classifier → Event Impact Score", GRN),
        ("L3 · PRESCRIBE", "manpower / barricades / diversions + city-wide officer optimization (ILP)", AMB),
        ("L4 · LEARN", "predicted-vs-actual playbook · model-drift detection · retrain triggers", RED)]
y, rh = Inches(1.55), Inches(1.04)
for title, body, c in rows:
    rect(s, Inches(0.6), y, Inches(3.5), Inches(0.9), c, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, Inches(0.6), y + Inches(0.26), Inches(3.5), Inches(0.5), [[(title, 14.5, TXT, True)]],
         align=PP_ALIGN.CENTER)
    rect(s, Inches(4.3), y, Inches(8.43), Inches(0.9), CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, Inches(4.6), y, Inches(8.0), Inches(0.9), [[(body, 13, TXT, False)]], anchor=MSO_ANCHOR.MIDDLE)
    y = y + rh
footer(s)

# =========================================================== 6 EIS
s = slide(); header(s, "L2 · THE SPINE", "Event Impact Score — engineering the missing target")
text(s, Inches(0.5), Inches(1.5), Inches(7.4), Inches(5.2),
     [[("Framed as expected vehicle-hours of delay:", 15, TXT, True)],
      [("EIS ≈ Clearance × Spatial reach × Baseline flow × Closure severity", 14, YEL, True)],
      [("", 8, TXT, False)],
      [("•  Clearance time is predicted (leakage-safe), not observed", 14, TXT, False)],
      [("•  Robust percentile-rank scaling → Low / Medium / High / Critical", 14, TXT, False)],
      [("•  VALIDATED against observed severity:", 14, TXT, True)],
      [("     Spearman 0.70  ·  AUC 0.89 flagging top-10% impact events", 14, GRN, True)]],
     anchor=MSO_ANCHOR.MIDDLE)
pic_fit(s, "eis_by_cause.png", Inches(7.95), Inches(1.5), Inches(5.0), Inches(5.25))
footer(s)

# =========================================================== 7 RIGOR
s = slide(); header(s, "MODELLING RIGOUR (KAGGLE-GRADE)", "10-model zoo · k-fold CV · leakage hunt · Optuna")
pic_fit(s, "closure_leaderboard.png", Inches(0.4), Inches(1.5), Inches(7.3), Inches(5.25))
text(s, Inches(7.95), Inches(1.5), Inches(4.9), Inches(5.25),
     [[("Headline results", 16, YEL, True)],
      [("Clearance time   MAE 2.88 h (blended)", 13.5, TXT, False)],
      [("Road closure     PR-AUC 0.471 · +469% vs baseline", 13.5, TXT, False)],
      [("EIS validity     Spearman 0.70 · AUC 0.89", 13.5, TXT, False)],
      [("Event forecast   +16.7% over lag-7 baseline", 13.5, TXT, False)],
      [("", 8, TXT, False)],
      [("Discipline that wins", 16, ABLUE, True)],
      [("•  Time-split CV + spatial holdout (no leakage)", 12.5, MUT, False)],
      [("•  Leakage hunt: dropped end-coords leaking @ AUC 0.976", 12.5, MUT, False)],
      [("•  Priority shown ~location-circular — reported honestly", 12.5, MUT, False)],
      [("•  Stacked ensemble + Optuna tuning", 12.5, MUT, False)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s)

# =========================================================== 8 FORECAST
s = slide(); header(s, "L1 · FORECAST", "Where, when & how many events — ahead of time")
pic_fit(s, "daily_volume.png", Inches(0.4), Inches(1.55), Inches(7.85), Inches(5.2))
text(s, Inches(8.45), Inches(1.55), Inches(4.4), Inches(5.2),
     [[("LightGBM-Poisson", 15.5, ABLUE, True)],
      [("corridor × day event counts — 16.7% better than the lag-7 baseline.", 13.5, TXT, False)],
      [("", 7, TXT, False)],
      [("Empirical-Bayes hotspots", 15.5, ABLUE, True)],
      [("recurring corridor × hour risk surface for pre-positioning.", 13.5, TXT, False)],
      [("", 7, TXT, False)],
      [("Hawkes self-excitation", 15.5, ABLUE, True)],
      [("accidents cluster — branching ratio 0.31 — quantifying contagion.", 13.5, TXT, False)]],
     anchor=MSO_ANCHOR.MIDDLE)
footer(s)

# =========================================================== 9 PRESCRIBE
s = slide(); header(s, "L3 · PRESCRIBE + OPTIMIZE", "The deliverable most teams never reach")
text(s, Inches(0.5), Inches(1.5), Inches(6.1), Inches(5.25),
     [[("Per-event plan", 16, AMB, True)],
      [("Officers, barricade units and a real road-network diversion (OSRM) — driven by the "
        "validated EIS tier + closure probability.", 13.5, TXT, False)],
      [("", 8, TXT, False)],
      [("City-wide optimization", 16, AMB, True)],
      [("Finite officer pool allocated by ILP (PuLP) — two doctrines:", 13.5, TXT, False)],
      [("•  Priority — cover highest-impact first (triage)", 13, TXT, False)],
      [("•  Efficiency — maximize total mitigated impact", 13, TXT, False)]], anchor=MSO_ANCHOR.MIDDLE)
rect(s, Inches(7.0), Inches(1.55), Inches(5.83), Inches(5.05), CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
rect(s, Inches(7.0), Inches(1.55), Inches(5.83), Inches(0.78), RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(7.25), Inches(1.55), Inches(5.4), Inches(0.78), [[("Surge day — 7 Mar 2024", 16, TXT, True)]],
     anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(7.35), Inches(2.6), Inches(5.2), Inches(3.9),
     [[("214 events in one day", 24, TXT, True)],
      [("≈ 1,096 officers needed", 19, RED, True)],
      [("vs ~120 available  →  a quantified deficit", 14, TXT, False)],
      [("", 10, TXT, False)],
      [("The system tells the commander exactly which events to staff first and how many more "
        "officers to escalate for — turning guesswork into math.", 13.5, MUT, False)]])
footer(s)

# =========================================================== 10 LEARN
s = slide(); header(s, "L4 · LEARN", "The post-event loop the brief explicitly asks for")
pic_fit(s, "learning_bias.png", Inches(0.4), Inches(1.5), Inches(7.7), Inches(5.25))
text(s, Inches(8.3), Inches(1.5), Inches(4.55), Inches(5.25),
     [[("After events close,", 15.5, RED, True)],
      [("we compare leak-free predictions with reality and build a playbook keyed by cause × H3 hex.", 13.5, TXT, False)],
      [("", 8, TXT, False)],
      [("It self-diagnoses:", 15.5, RED, True)],
      [("the model under-predicts infrastructure events — potholes by +8.0 h, water-logging +7.1 h — "
        "and flags those cells for retraining.", 13.5, TXT, False)],
      [("", 8, TXT, False)],
      [("Closes the “no post-event learning” gap most teams ignore.", 13.5, ABLUE, True)]],
     anchor=MSO_ANCHOR.MIDDLE)
footer(s)

# =========================================================== 11 PRODUCT
s = slide(); header(s, "THE PRODUCT", "A live decision-support dashboard")
tabs = [("🎯  Event Simulator", "score any what-if event → EIS, clearance P50/P90, full deployment plan"),
        ("🗺️  City Risk Map", "all 8,173 events on a Bengaluru map, coloured by impact tier"),
        ("📅  Surge-Day Replay", "replay a real day, set an officer budget → live allocation + deficit"),
        ("📈  Forecast & Hotspots", "corridor × hour risk heatmap, Hawkes, the 2-AM explainer"),
        ("🔁  Learning Loop", "predicted-vs-actual playbook + drift alerts")]
y = Inches(1.55)
for t1, t2 in tabs:
    rect(s, Inches(0.6), y, Inches(3.6), Inches(0.82), BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, Inches(0.78), y, Inches(3.4), Inches(0.82), [[(t1, 13.5, TXT, True)]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, Inches(4.35), y, Inches(8.4), Inches(0.82), CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, Inches(4.6), y, Inches(8.0), Inches(0.82), [[(t2, 13, TXT, False)]], anchor=MSO_ANCHOR.MIDDLE)
    y = y + Inches(0.95)
text(s, Inches(0.6), Inches(6.4), Inches(12.2), Inches(0.55),
     [[("Live proof: the same breakdown scores EIS 17 at 2 AM vs 78 at 9 AM rush; a VIP movement "
        "scores 99.7 (Critical) with a real OSRM diversion drawn on the map.", 12.5, MUT, False)]])
footer(s)

# =========================================================== 12 WHY WIN
s = slide(); header(s, "WHY THIS WINS", "Seven things the field won't have together")
items = [("Correct reframe", "incident log, not congestion data", BLUE),
         ("Engineered + validated EIS", "Spearman 0.70 — the missing target, proven", GRN),
         ("k-fold model zoo", "10 models, stacking, Optuna — not one model", AMB),
         ("Prescription + ILP", "we allocate scarce officers; others only predict", RED),
         ("Post-event learning", "L4 playbook + drift — the 3rd gap, solved", ABLUE),
         ("Real-ML simulator", "runs the trained models, not random numbers", GRN),
         ("Honesty & rigor", "leakage hunt, time-split CV, stated limits", AMB)]
xs = [Inches(0.6), Inches(4.85), Inches(9.1)]
for i, (t1, t2, c) in enumerate(items):
    x = xs[i % 3]; y = Inches(1.6) + (i // 3) * Inches(1.74)
    card(s, x, y, Inches(3.95), Inches(1.56), c, t1, t2)
footer(s)

# =========================================================== 13 IMPACT
s = slide(); header(s, "IMPACT, ROADMAP & HONESTY", "Operational value — and what we'd add next")
text(s, Inches(0.5), Inches(1.55), Inches(6.0), Inches(5.2),
     [[("Operational value", 16, GRN, True)],
      [("•  Quantified impact before the event, not after", 13.5, TXT, False)],
      [("•  Optimal, defensible resource deployment", 13.5, TXT, False)],
      [("•  Faster clearance via right-sized response", 13.5, TXT, False)],
      [("•  A city that gets smarter after every event", 13.5, TXT, False)],
      [("", 10, TXT, False)],
      [("Roadmap", 16, ABLUE, True)],
      [("•  Plug in live flow (Google/TomTom) to calibrate EIS", 13.5, TXT, False)],
      [("•  Spatio-temporal GNN forecasting · officer-shift scheduling", 13.5, TXT, False)]],
     anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(6.85), Inches(1.55), Inches(5.95), Inches(5.2),
     [[("Honest limitations (we say so)", 16, RED, True)],
      [("•  No ground-truth flow → EIS validated vs proxies", 13.5, TXT, False)],
      [("•  Manpower data sparse (1.6%) → optimization + rules,", 13.5, TXT, False)],
      [("    not over-claimed as learned", 13.5, TXT, False)],
      [("•  5 months of data → calendar priors mitigate seasonality", 13.5, TXT, False)],
      [("", 12, TXT, False)],
      [("Stack", 16, ABLUE, True)],
      [("Python · LightGBM/XGBoost/CatBoost · scikit-learn · Optuna ·", 12.5, MUT, False)],
      [("H3 · PuLP · OSRM · Streamlit + pydeck", 12.5, MUT, False)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s)

# =========================================================== 14 CLOSE
s = slide()
rect(s, 0, 0, W, Inches(0.22), YEL)
rect(s, 0, Inches(7.28), W, Inches(0.22), BLUE)
text(s, Inches(0.9), Inches(2.3), Inches(11.6), Inches(1.2), [[("From reaction to anticipation.", 40, TXT, True)]])
text(s, Inches(0.95), Inches(3.7), Inches(11.6), Inches(1.3),
     [[("Forecast the event · quantify the impact · deploy the optimal response · "
        "learn from every outcome.", 18, YEL, False)]])
rect(s, Inches(0.95), Inches(5.4), Inches(4.2), Pt(2), LINE)
text(s, Inches(0.95), Inches(5.6), Inches(11.6), Inches(1.5),
     [[(TEAM_NAME, 17, TXT, True)],
      [(TEAM_MEMBERS, 13.5, YEL, False)],
      [("GRIDLOCK — Event-Driven Congestion Intelligence  ·  Flipkart Gridlock Hackathon 2.0 · Round 2",
        12.5, MUT, False)]])

out = C.ROOT / "Gridlock_Pitch_Deck.pptx"
prs.save(str(out))
print(f"Saved dark deck -> {out}  ({len(list(prs.slides))} slides)")
